"""Independent CONV-007 export assertions. Usage: python scripts/verify-table-workbooks.py DIR.
Read only freshly generated synthetic test outputs, never an untrusted uploaded workbook.
"""
import argparse
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook

parser=argparse.ArgumentParser(description=__doc__)
parser.add_argument('directory',type=Path)
parser.add_argument('--stems',nargs='+',choices=['inventory','literal','ocr'],default=['inventory','literal','ocr'],help='Explicit subset of native exports; default requires all six workbooks.')
args=parser.parse_args()
directory=args.directory
ns={'table':'urn:oasis:names:tc:opendocument:xmlns:table:1.0','text':'urn:oasis:names:tc:opendocument:xmlns:text:1.0','office':'urn:oasis:names:tc:opendocument:xmlns:office:1.0'}
expected={
    'inventory':[['Code','Item','Quantity','Amount'],['00127','Notebook',5,12.5],['00128','Pen pack',12,30],['00129','Folder',0,0],['00130','Credit',2,-5.25],['00131','Desk lamp',1,100]],
    'literal':[['Code','Label','Rate','Note'],['0007','Standard','12,5%','=1+1'],['0008','Reduced','5,0%',None],['0009','Amount','1.234,50','Text, not a formula'],['0010','Large ID','0,0%','1234567890123456']],
    'ocr':[['Product','Units','Total'],['Notebook',5,12.5],['Pen pack',12,30],['Desk lamp',1,100]],
}
for stem,rows in expected.items():
    if stem not in args.stems:continue
    for extension in ['xlsx','ods']:
        path=directory/f'{stem}.{extension}'
        assert path.exists(), f'Missing expected export: {path.name}'
        with ZipFile(path) as archive:
            assert not any('externallink' in name.lower() or 'vbaproject' in name.lower() for name in archive.namelist())
            if extension=='xlsx':
                workbook=load_workbook(path,data_only=False)
                assert workbook.sheetnames==['Extracted table']
                sheet=workbook.active
                actual=list(map(list,sheet.values))
                assert all(cell.data_type!='f' and not cell.hyperlink for row in sheet for cell in row)
                assert sheet.auto_filter.ref == f'A1:{chr(64+len(rows[0]))}{len(rows)}'
                if stem!='ocr':assert sheet['A2'].data_type=='s'
                assert sheet.cell(2,3 if stem=='ocr' else 4 if stem=='inventory' else 3).data_type==('s' if stem=='literal' else 'n')
            else:
                first=archive.infolist()[0]
                assert first.filename=='mimetype' and first.compress_type==0 and not first.extra
                assert archive.read('mimetype')==b'application/vnd.oasis.opendocument.spreadsheet'
                root=ET.fromstring(archive.read('content.xml'))
                actual=[]
                for row in root.findall('.//table:table-row',ns):
                    values=[]
                    for cell in row.findall('table:table-cell',ns):
                        assert f'{{{ns["table"]}}}formula' not in cell.attrib
                        kind=cell.get(f'{{{ns["office"]}}}value-type')
                        value=float(cell.get(f'{{{ns["office"]}}}value')) if kind=='float' else '\n'.join(''.join(p.itertext()) for p in cell.findall('text:p',ns)) or None
                        repeat=int(cell.get(f'{{{ns["table"]}}}number-columns-repeated','1'))
                        values.extend([value]*repeat)
                    actual.append(values)
            assert actual==rows,(path.name,actual,rows)
        if stem=='inventory':
            assert sum(row[2] for row in actual[1:])==20
            assert abs(sum(row[3] for row in actual[1:])-137.25)<1e-9
            assert abs(sum(row[3] for row in actual[1:])/5-27.45)<1e-9
        print(f'{path.name}: independent cells, types, numeric totals and literal-text checks passed')
